"""Fill the GitHub repository URL (and optionally the team name) into every submission file.

    python scripts/finalize_submission.py --repo-url https://github.com/<user>/kendali
    python scripts/finalize_submission.py --repo-url https://github.com/<user>/kendali --team "Kendali"
    python scripts/finalize_submission.py --check      # list the placeholders that are still open
    python scripts/finalize_submission.py --repo-url https://github.com/<user>/kendali --pdf   # + PDFs in submission/pdf/

Touches README.md, docs/*.md, submission/*.md and the pitch deck (slide text and speaker notes).
Only exact placeholders are replaced, so running it twice is harmless.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
REPO_PLACEHOLDER = "github.com/<github-user>/kendali"
TEAM_PLACEHOLDER = "Kendali"  # current team name in all files
DECK = ROOT / "submission" / "Kendali_FinTechathon2026_Deck.pptx"


def text_files() -> list[Path]:
    files = [ROOT / "README.md"]
    files += sorted((ROOT / "docs").glob("*.md"))
    files += sorted((ROOT / "submission").glob("*.md"))
    return [f for f in files if f.exists()]


def deck_paragraphs(prs):
    for slide in prs.slides:
        frames = [sh.text_frame for sh in slide.shapes if sh.has_text_frame]
        if slide.has_notes_slide:
            frames.append(slide.notes_slide.notes_text_frame)
        for tf in frames:
            yield from tf.paragraphs


def replace_in_deck(pairs: list[tuple[str, str]], dry: bool) -> int:
    from pptx import Presentation

    prs = Presentation(str(DECK))
    hits = 0
    for para in deck_paragraphs(prs):
        for run in para.runs:
            new = run.text
            for a, b in pairs:
                new = new.replace(a, b)
            if new != run.text:
                hits += 1
                if not dry:
                    run.text = new
    if hits and not dry:
        prs.save(str(DECK))
    return hits


def check() -> int:
    open_items = 0
    for f in text_files():
        for i, line in enumerate(f.read_text(encoding="utf-8").splitlines(), 1):
            if "<github-user>" in line or "[TEAM NAME]" in line or "<repo-url>" in line:
                print(f"{f.relative_to(ROOT)}:{i}: {line.strip()[:110]}")
                open_items += 1
    from pptx import Presentation

    for para in deck_paragraphs(Presentation(str(DECK))):
        t = "".join(r.text for r in para.runs)
        if "<github-user>" in t or "[TEAM NAME]" in t or "[team]" in t:
            print(f"{DECK.relative_to(ROOT)}: {t.strip()[:110]}")
            open_items += 1
    print(f"{open_items} open placeholder(s)")
    return open_items


PDF_JOBS = [  # (output name, markdown sources) -> submission/pdf/
    ("Kendali_TechnicalDocumentation.pdf", ["docs/TECHNICAL_DOCUMENTATION.md"]),
    ("Kendali_SecuritySelfAssessment.pdf", ["docs/SECURITY_SELF_ASSESSMENT.md"]),
    ("Kendali_ExecutionEvidence.pdf", ["docs/EXECUTION_EVIDENCE.md", "reports/TESTNET_EVIDENCE.md", "reports/EVAL_REPORT.md"]),
]
CSS = "body{font-family:sans-serif;font-size:10pt} table{border-collapse:collapse} td,th{border:1px solid #999;padding:2px 4px;font-size:8pt} code{font-size:8pt} img{max-width:100%}"


def export_pdfs() -> None:
    """Render the Markdown documents to PDF email attachments (needs `pip install -e .[docs]`)."""
    from markdown_pdf import MarkdownPdf, Section

    out = ROOT / "submission" / "pdf"
    out.mkdir(parents=True, exist_ok=True)
    for name, sources in PDF_JOBS:
        pdf = MarkdownPdf(toc_level=2, optimize=True)
        for src in sources:
            path = ROOT / src
            text = re.sub(r"<!--.*?-->", "", path.read_text(encoding="utf-8"), flags=re.S)
            pdf.add_section(Section(text, root=str(path.parent), toc=False), user_css=CSS)
        pdf.meta["title"] = name.removesuffix(".pdf").replace("_", " ")
        pdf.meta["author"] = "Team Kendali"
        pdf.save(str(out / name))
        print("wrote", (out / name).relative_to(ROOT))


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo-url", help="e.g. https://github.com/<user>/kendali")
    ap.add_argument("--team", help="team name to use instead of 'Kendali' in the team lines")
    ap.add_argument("--check", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--pdf", action="store_true", help="also export the documents to submission/pdf/")
    a = ap.parse_args()
    if a.pdf and not a.repo_url:
        export_pdfs()
        return
    if a.check or not a.repo_url:
        sys.exit(1 if check() else 0)

    url = a.repo_url.strip().rstrip("/")
    url = re.sub(r"\.git$", "", url)
    m = re.match(r"^https://github\.com/([A-Za-z0-9-]+)/([A-Za-z0-9._-]+)$", url)
    if not m:
        sys.exit("--repo-url must look like https://github.com/<user>/<repo>")
    bare = url.removeprefix("https://")

    pairs = [(REPO_PLACEHOLDER, bare)]
    team_pairs = []
    if a.team and a.team != TEAM_PLACEHOLDER:
        team_pairs = [(f"Team {TEAM_PLACEHOLDER}", f"Team {a.team}"), (f"team {TEAM_PLACEHOLDER}", f"team {a.team}"),
                      (f"**Team:** {TEAM_PLACEHOLDER}", f"**Team:** {a.team}"), (f"| Team | **{TEAM_PLACEHOLDER}**", f"| Team | **{a.team}**")]

    for f in text_files():
        s = f.read_text(encoding="utf-8")
        new = s
        for x, y in pairs + team_pairs:
            new = new.replace(x, y)
        # drop the reminder comments once the URL is set
        new = re.sub(r"\s*<!-- SET-REPO-URL[^>]*-->", "", new) if REPO_PLACEHOLDER not in new else new
        if new != s:
            print(("would update " if a.dry_run else "updated ") + str(f.relative_to(ROOT)))
            if not a.dry_run:
                f.write_text(new, encoding="utf-8")
    n = replace_in_deck(pairs + team_pairs, a.dry_run)
    print(f"deck: {n} text run(s) {'would change' if a.dry_run else 'changed'}")
    check()
    if a.pdf and not a.dry_run:
        export_pdfs()


if __name__ == "__main__":
    main()
